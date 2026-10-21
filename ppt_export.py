"""
Step: export a structured Chapter into a downloadable .pptx, one slide per
section, so teachers can use it in class without the video player.

Implements issue #1 exactly:
    def export_chapter_to_pptx(chapter, out_path: str) -> str

Input : a Chapter object (schema.py) — chapter.sections is a list, each
        section has .title, .hook, .concept, .indian_example, .key_points,
        .misconception
Output: a .pptx file saved to out_path. Returns out_path (matches the
        return-a-path convention used by slides.py / video_builder.py).

No Gemini, TTS, or video code is touched. Colors are pulled from
slides.SUBJECT_THEMES so the deck visually matches the generated video.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from pptx.oxml.ns import qn
from lxml import etree

from slides import SUBJECT_THEMES, LABELS, _pretty_equation

# Complex-script (Kannada) font. Nirmala UI ships with Windows 10+, which is what most
# school PCs run; other apps substitute a Kannada-capable font automatically.
KN_FONT = "Nirmala UI"


def _kn(run, language):
    """Mark a run's complex-script font so PowerPoint shapes Kannada properly."""
    if language != "Kannada":
        return
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:cs", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", KN_FONT)
    run.font.name = KN_FONT


BLANK_LAYOUT_INDEX = 6  # "Blank" layout in the default pptx template
TITLE_ONLY_LAYOUT_INDEX = 5  # "Title Only" layout


def _theme_colors(subject: str):
    theme = SUBJECT_THEMES.get(subject, SUBJECT_THEMES["Other"])
    return {
        "bg": RGBColor(*theme["bg"]),
        "accent": RGBColor(*theme["accent"]),
        "text": RGBColor(*theme["text"]),
    }


def _set_background(slide, rgb: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def _add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def _add_title_slide(prs: Presentation, chapter_title: str, subject: str, language: str = "English"):
    colors = _theme_colors(subject)
    layout = prs.slide_layouts[BLANK_LAYOUT_INDEX]
    slide = prs.slides.add_slide(layout)
    _set_background(slide, colors["bg"])

    tf = _add_textbox(slide, Inches(0.7), Inches(2.6), Inches(8.6), Inches(2.2))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = chapter_title
    _kn(run, language)
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = colors["accent"]

    sub = tf.add_paragraph()
    sub.alignment = PP_ALIGN.CENTER
    sub_run = sub.add_run()
    sub_run.text = LABELS.get(language, LABELS["English"])["subjects"].get(subject, subject) if language == "Kannada" else subject
    _kn(sub_run, language)
    sub_run.font.size = Pt(20)
    sub_run.font.color.rgb = colors["text"]


def _add_section_slide(prs: Presentation, section, subject: str, language: str = "English"):
    lab = LABELS.get(language, LABELS["English"])
    colors = _theme_colors(subject)
    layout = prs.slide_layouts[BLANK_LAYOUT_INDEX]
    slide = prs.slides.add_slide(layout)
    _set_background(slide, colors["bg"])

    # Title = section.title
    title_tf = _add_textbox(slide, Inches(0.5), Inches(0.35), Inches(9), Inches(0.9))
    title_p = title_tf.paragraphs[0]
    title_run = title_p.add_run()
    title_run.text = section.title
    _kn(title_run, language)
    title_run.font.size = Pt(30)
    title_run.font.bold = True
    title_run.font.color.rgb = colors["accent"]

    # Body = concept + indian_example
    body_tf = _add_textbox(slide, Inches(0.5), Inches(1.3), Inches(9), Inches(2.3))
    concept_p = body_tf.paragraphs[0]
    concept_run = concept_p.add_run()
    concept_run.text = section.concept
    _kn(concept_run, language)
    concept_run.font.size = Pt(18)
    concept_run.font.color.rgb = colors["text"]

    if getattr(section, "key_equation", None):
        eq_p = body_tf.add_paragraph()
        eq_p.space_before = Pt(8)
        eq_label = eq_p.add_run()
        eq_label.text = lab["equation"] + ": "
        eq_label.font.bold = True
        eq_label.font.size = Pt(16)
        eq_label.font.color.rgb = colors["accent"]
        _kn(eq_label, language)
        eq_run = eq_p.add_run()
        eq_run.text = _pretty_equation(section.key_equation, subject)
        eq_run.font.size = Pt(16)
        eq_run.font.bold = True
        eq_run.font.color.rgb = colors["text"]

    example_p = body_tf.add_paragraph()
    example_p.space_before = Pt(12)
    example_label = example_p.add_run()
    example_label.text = lab["example"] + ": "
    _kn(example_label, language)
    example_label.font.bold = True
    example_label.font.size = Pt(16)
    example_label.font.color.rgb = colors["accent"]

    example_run = example_p.add_run()
    example_run.text = section.indian_example
    _kn(example_run, language)
    example_run.font.size = Pt(16)
    example_run.font.color.rgb = colors["text"]

    # Bullet list = key_points
    points_tf = _add_textbox(slide, Inches(0.5), Inches(3.9), Inches(9), Inches(3.2))
    for i, point in enumerate(section.key_points):
        p = points_tf.paragraphs[0] if i == 0 else points_tf.add_paragraph()
        run = p.add_run()
        run.text = f"\u2022 {point}"
        _kn(run, language)
        run.font.size = Pt(16)
        run.font.color.rgb = colors["text"]
        p.space_after = Pt(8)


def export_chapter_to_pptx(chapter, out_path: str, language: str = "English") -> str:
    """
    Build a .pptx from a Chapter object: one title slide + one slide per
    section (title / concept + indian_example / key_points bullets).

    Returns out_path.
    """
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    _add_title_slide(prs, chapter.chapter_title, chapter.subject, language)

    for section in chapter.sections:
        _add_section_slide(prs, section, chapter.subject, language)

    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    # Quick manual smoke test with dummy data matching schema.py's shape.
    from schema import Chapter, Section, Quiz

    dummy_chapter = Chapter(
        chapter_title="Chemical Reactions and Equations",
        subject="Science",
        sections=[
            Section(
                title="What is a Chemical Reaction?",
                hook="Ever noticed rust on an old gate?",
                concept="A chemical reaction is a process where substances "
                        "change into new substances with different properties.",
                indian_example="Rusting of iron gates and railings during "
                                "the monsoon in Indian households.",
                key_points=["Reactants become products", "New substances form",
                            "Energy is absorbed or released"],
                misconception="Mixing two liquids always means a reaction happened.",
                narration="Placeholder narration text for testing purposes only.",
                quiz=Quiz(
                    question="What forms when iron reacts with oxygen and moisture?",
                    options=["Rust", "Salt", "Gas", "Water"],
                    answer_index=0,
                    explanation="Iron reacts with oxygen and moisture to form rust (iron oxide).",
                ),
            )
        ],
    )
    export_chapter_to_pptx(dummy_chapter, "test_output.pptx")
    print("Saved test_output.pptx")